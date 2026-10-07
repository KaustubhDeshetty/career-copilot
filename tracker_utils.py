"""
tracker_utils.py
Saves pipeline results to an Excel tracker (applications.xlsx).

Two sheets in one workbook:
  "Applications"      -- Apply-mode results (company, role, fit, verdict)
  "Interview Sessions" -- Interview-mode practice log (topic, difficulty, score)

Deliberately zero-cost: everything written here is extracted from data the
pipeline already produced, via simple text parsing -- no additional LLM
calls, since token budget is already tight on this account tier.
"""

import os
import re
from datetime import datetime
from openpyxl import Workbook, load_workbook

TRACKER_PATH = "./applications.xlsx"

APPLICATIONS_SHEET = "Applications"
INTERVIEW_SHEET = "Interview Sessions"

TRACKER_HEADERS = ["Date", "Company", "Role", "Fit Score", "Recommendation", "Notes"]
INTERVIEW_HEADERS = ["Date", "Topic", "Difficulty", "Questions", "Critic Verdict", "Rounds", "Notes"]


def extract_company_role(job_posting_text: str) -> tuple[str, str]:
    """Best-effort extraction of role and company from the first couple
    lines of a job posting. Most postings lead with 'Role Title' then
    'Company Name' before the body text starts (e.g. the Meridian AI JD
    format used in testing). Falls back gracefully on postings that don't
    follow this shape -- worst case, edit the cell directly in Excel."""
    lines = [l.strip() for l in job_posting_text.strip().split("\n") if l.strip()]
    role = lines[0] if len(lines) > 0 else "Unknown role"
    company = lines[1] if len(lines) > 1 and len(lines[1]) < 60 else "Unknown company"
    return role, company


def extract_fit_summary(draft: str) -> tuple[str, str]:
    """Pulls the fit ratio and recommendation out of an Apply-mode draft's
    Fit Assessment section. Counts MET/PARTIALLY MET/NOT MET directly from
    the table's bolded status markers -- more robust than matching one
    exact summary-sentence phrasing, since the model phrases that sentence
    differently once PARTIALLY MET rows are involved (not just binary
    MET/NOT MET, which is all the phrase-matching version could handle)."""
    if not draft:
        return "Unknown", "Unknown"

    # Tier 1 -- bolded status markers. Unchanged: this is what the model
    # actually produces most of the time, and it can't false-positive on
    # prose. Note "**MET**" cannot match inside "**NOT MET**", since the
    # characters before MET there are "T " rather than "**".
    not_met = len(re.findall(r"\*\*NOT MET\*\*", draft, re.IGNORECASE))
    partially_met = len(re.findall(r"\*\*PARTIALLY MET\*\*", draft, re.IGNORECASE))
    fully_met = len(re.findall(r"\*\*MET\*\*", draft, re.IGNORECASE))

    # Tier 2 -- the model didn't bold this time. Nothing in the Writer prompt
    # actually REQUIRES bold, so this is a normal outcome, not a rare one.
    # Match unbolded markers, but only where one sits in a table cell or after
    # a ':'/bullet -- never loose in prose, or "the candidate has MET the bar"
    # and words like "METRICS" start counting as requirements.
    if fully_met + partially_met + not_met == 0:
        status = r"(?:\*\*|__)?(NOT MET|PARTIALLY MET|MET)(?:\*\*|__)?"
        counts = {"MET": 0, "PARTIALLY MET": 0, "NOT MET": 0}
        for line in draft.splitlines():
            hit = re.search(rf"(?:\||:|[-*])\s*{status}\s*(?:\||$|\.)", line, re.IGNORECASE)
            if hit:
                counts[hit.group(1).upper()] += 1
        fully_met, partially_met, not_met = counts["MET"], counts["PARTIALLY MET"], counts["NOT MET"]

    total = fully_met + partially_met + not_met
    if total > 0:
        fit_score = f"{fully_met} of {total} core requirements met"
        if partially_met > 0:
            fit_score += f" ({partially_met} partially met)"
    else:
        # Tier 3 -- no table at all, but the Writer prompt asks for a summary
        # line phrased "3 of 4 core requirements met". The previous version of
        # this fallback required a literal "Overall fit:" prefix that the
        # prompt never asks for, so it essentially never fired -- which is why
        # both existing rows in applications.xlsx read "Unknown".
        fit_match = re.search(r"(\d+)\s+of\s+(\d+)\s+core requirements met", draft, re.IGNORECASE)
        fit_score = f"{fit_match.group(1)} of {fit_match.group(2)} core requirements met" if fit_match else "Unknown"

    rec_match = re.search(r"Recommendation:?\**\s*\**\s*([^\n]+)", draft, re.IGNORECASE)
    recommendation = rec_match.group(1).strip() if rec_match else "Unknown"
    # Split on a real sentence boundary (period + space + capital letter),
    # not just any period -- a plain ".".split() breaks abbreviations like
    # "e.g." mid-word, cutting off text right after "(e".
    recommendation = re.split(r"\.\s+(?=[A-Z])", recommendation)[0]
    recommendation = recommendation.replace("**", "").strip(" *")

    return fit_score, recommendation


def _open_workbook() -> Workbook:
    """Loads the tracker, or builds an empty workbook with no sheets at all.

    A fresh Workbook() ships with one sheet already titled "Sheet"; dropping
    it here means every sheet in this file is created explicitly by name in
    _get_sheet, so the workbook never ends up with a stray default tab."""
    if os.path.exists(TRACKER_PATH):
        return load_workbook(TRACKER_PATH)
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _get_sheet(wb: Workbook, title: str, headers: list):
    """Fetches a sheet BY NAME, creating it with its header row if absent.

    Addressing sheets by name rather than wb.active matters now that the
    workbook holds more than one: wb.active returns whichever tab was
    focused when the file was last saved, so the previous version would
    have started appending applications onto the interview sheet (or vice
    versa) as soon as a second sheet existed."""
    if title in wb.sheetnames:
        return wb[title]
    ws = wb.create_sheet(title)
    ws.append(headers)
    return ws


def _append_row(sheet_title: str, headers: list, row: list) -> None:
    wb = _open_workbook()
    ws = _get_sheet(wb, sheet_title, headers)
    ws.append(row)
    wb.save(TRACKER_PATH)


def save_application(company: str, role: str, fit_score: str, recommendation: str, notes: str = "") -> None:
    """Appends one row to the Applications sheet, creating file/sheet if needed."""
    _append_row(APPLICATIONS_SHEET, TRACKER_HEADERS, [
        datetime.now().strftime("%Y-%m-%d"),
        company,
        role,
        fit_score,
        recommendation,
        notes,
    ])


def save_interview_session(topic: str, difficulty: str, num_questions: int,
                           verdict: str, rounds: int, notes: str = "") -> None:
    """Appends one row to the Interview Sessions sheet -- the practice log
    counterpart to save_application. Same zero-LLM-cost principle: topic and
    difficulty come straight off the Router's extracted state, the question
    count from the parsed draft, the verdict from the Critic."""
    _append_row(INTERVIEW_SHEET, INTERVIEW_HEADERS, [
        datetime.now().strftime("%Y-%m-%d"),
        topic,
        difficulty,
        num_questions,
        verdict,
        rounds,
        notes,
    ])


def load_sheet(sheet_title: str) -> tuple[list, list]:
    """Reads one sheet back as (headers, rows) so the UI can display what's
    been tracked, rather than the file being a write-only side effect.

    Returns empty lists when the file or sheet doesn't exist yet, so callers
    can render an honest 'nothing logged yet' state instead of erroring."""
    if not os.path.exists(TRACKER_PATH):
        return [], []
    wb = load_workbook(TRACKER_PATH, read_only=True)
    try:
        if sheet_title not in wb.sheetnames:
            return [], []
        rows = [list(r) for r in wb[sheet_title].iter_rows(values_only=True)]
    finally:
        # read_only mode holds the file open until closed -- on Windows that
        # would keep a lock on applications.xlsx for the life of the app.
        wb.close()
    if not rows:
        return [], []
    return rows[0], [r for r in rows[1:] if any(c is not None and str(c).strip() for c in r)]


def load_applications() -> tuple[list, list]:
    return load_sheet(APPLICATIONS_SHEET)


def load_interview_sessions() -> tuple[list, list]:
    return load_sheet(INTERVIEW_SHEET)


# ---------------------------------------------------------------------------
# Duplicate detection
#
# app.py's original guard lived only in st.session_state, which Streamlit
# clears on a page refresh -- so refresh-then-click-Save wrote a second
# identical row, which is how applications.xlsx picked up a duplicate pair in
# the first place. Checking the workbook itself makes the guard survive a
# refresh, a browser restart, and a whole new session.
#
# Scope is deliberately "same entry, same day": re-running Apply for the same
# role next week is a legitimate new row (the fit may have changed after a
# resume edit), so the date is part of the identity rather than matching on
# company+role alone and permanently locking out a re-run.
# ---------------------------------------------------------------------------

def _norm(value) -> str:
    """Compares cells the way a human would: ignoring case, surrounding space,
    and the difference between an empty cell and a missing one."""
    return "" if value is None else str(value).strip().lower()


def _today() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def row_exists(sheet_title: str, match: dict) -> bool:
    """True if any row in the sheet matches every column/value pair in `match`.

    Unknown column names return False rather than raising -- a caller asking
    about a column this sheet doesn't have cannot have written that row, and a
    duplicate check is not worth crashing a Save button over."""
    headers, rows = load_sheet(sheet_title)
    if not headers or not rows:
        return False

    try:
        indexes = {headers.index(col): _norm(val) for col, val in match.items()}
    except ValueError:
        return False

    return any(
        all(_norm(row[i]) == want for i, want in indexes.items() if i < len(row))
        for row in rows
    )


def application_exists(company: str, role: str, on_date: str | None = None) -> bool:
    """Has this company/role already been logged today?"""
    return row_exists(APPLICATIONS_SHEET, {
        "Date": on_date or _today(),
        "Company": company,
        "Role": role,
    })


def interview_session_exists(topic: str, difficulty: str, on_date: str | None = None) -> bool:
    """Has this topic/difficulty practice session already been logged today?"""
    return row_exists(INTERVIEW_SHEET, {
        "Date": on_date or _today(),
        "Topic": topic,
        "Difficulty": difficulty,
    })
